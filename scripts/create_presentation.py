"""
System Architecture Presentation - Mentor Review
EV Charging Recommendation System
Clean, professional design - Architecture focused
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as RgbColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Colors - clean professional palette
NAVY = RgbColor(0x1E, 0x3A, 0x5F)
BLUE = RgbColor(0x3B, 0x82, 0xF6)
TEAL = RgbColor(0x14, 0xB8, 0xA6)
ORANGE = RgbColor(0xF5, 0x73, 0x16)
RED = RgbColor(0xEF, 0x44, 0x44)
SLATE = RgbColor(0x64, 0x74, 0x8B)
LIGHT_BG = RgbColor(0xF8, 0xFA, 0xFC)
WHITE = RgbColor(0xFF, 0xFF, 0xFF)
PURPLE = RgbColor(0x8B, 0x5C, 0xF6)


def add_title_bar(slide, prs, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    t = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.6))
    tf = t.text_frame
    tf.paragraphs[0].text = title
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(9), Inches(0.3))
        tf = s.text_frame
        tf.paragraphs[0].text = subtitle
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = RgbColor(0xB4, 0xC4, 0xD4)


def arrow_down(slide, x, y, h=0.35):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(0.15), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = SLATE
    a.line.fill.background()


def arrow_right(slide, x, y, w=0.4):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.15))
    a.fill.solid()
    a.fill.fore_color.rgb = SLATE
    a.line.fill.background()


def box(slide, x, y, w, h, color, title, subtitle=None):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid()
    b.fill.fore_color.rgb = color
    b.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(x), Inches(y + (h - 14) / 2 / 14), Inches(w), Inches(h * 0.6))
    tf = tb.text_frame
    tf.paragraphs[0].text = title
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    if subtitle:
        stb = slide.shapes.add_textbox(Inches(x), Inches(y + h * 0.45), Inches(w), Inches(h * 0.4))
        tf = stb.text_frame
        tf.paragraphs[0].text = subtitle
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER


# ============================================================
# SLIDE 1: Title
# ============================================================
def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()

    dec1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1), Inches(-1), Inches(4), Inches(4))
    dec1.fill.solid()
    dec1.fill.fore_color.rgb = RgbColor(0x2D, 0x4A, 0x6F)
    dec1.line.fill.background()

    dec2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(3.5), Inches(4), Inches(4))
    dec2.fill.solid()
    dec2.fill.fore_color.rgb = RgbColor(0x2D, 0x4A, 0x6F)
    dec2.line.fill.background()

    t = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(1))
    tf = t.text_frame
    tf.paragraphs[0].text = "EV Charging Recommendation"
    tf.paragraphs[0].font.size = Pt(44)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE

    s = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(8.4), Inches(0.8))
    tf = s.text_frame
    tf.paragraphs[0].text = "System Architecture"
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.color.rgb = BLUE

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.6), Inches(3), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    m = slide.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(8.4), Inches(0.5))
    tf = m.text_frame
    tf.paragraphs[0].text = "Mentor Technical Review"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = RgbColor(0xB4, 0xC4, 0xD4)


# ============================================================
# SLIDE 2: High-Level Architecture
# ============================================================
def slide_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "System Overview", "Architecture layers")

    # Client
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(1.4), Inches(3), Inches(0.6))
    c.fill.solid()
    c.fill.fore_color.rgb = SLATE
    c.line.fill.background()
    ct = slide.shapes.add_textbox(Inches(3.5), Inches(1.5), Inches(3), Inches(0.4))
    tf = ct.text_frame
    tf.paragraphs[0].text = "Client Application"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    arrow_down(slide, 4.92, 2.0, 0.35)

    # Gateway
    g = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3), Inches(2.5), Inches(4), Inches(0.7))
    g.fill.solid()
    g.fill.fore_color.rgb = BLUE
    g.line.fill.background()
    gt = slide.shapes.add_textbox(Inches(3), Inches(2.65), Inches(4), Inches(0.4))
    tf = gt.text_frame
    tf.paragraphs[0].text = "FastAPI Gateway"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Services
    services = [
        (1.5, "Map Matching", TEAL),
        (4.2, "Demand Detection", ORANGE),
        (6.9, "Candidate Search", PURPLE),
    ]
    for x, name, color in services:
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.5), Inches(2.2), Inches(0.8))
        s.fill.solid()
        s.fill.fore_color.rgb = color
        s.line.fill.background()
        st = slide.shapes.add_textbox(Inches(x), Inches(3.6), Inches(2.2), Inches(0.6))
        tf = st.text_frame
        tf.paragraphs[0].text = name
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    for x in [2.6, 5.1, 7.6]:
        arrow_down(slide, x, 3.2, 0.3)

    # Data layer
    data = [
        (1.2, "GraphHopper", "Routing", ORANGE),
        (4.2, "PostgreSQL", "Snapshots", TEAL),
        (7.2, "Redis", "Cache + State", BLUE),
    ]
    for x, name, sub, color in data:
        d = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.7), Inches(2.4), Inches(0.8))
        d.fill.solid()
        d.fill.fore_color.rgb = color
        d.line.fill.background()
        dt = slide.shapes.add_textbox(Inches(x), Inches(4.8), Inches(2.4), Inches(0.35))
        tf = dt.text_frame
        tf.paragraphs[0].text = name
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        dst = slide.shapes.add_textbox(Inches(x), Inches(5.1), Inches(2.4), Inches(0.3))
        tf = dst.text_frame
        tf.paragraphs[0].text = sub
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    for x in [2.6, 5.1, 7.6]:
        arrow_down(slide, x, 4.3, 0.35)


# ============================================================
# SLIDE 3: Request Flow
# ============================================================
def slide_flow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "Request Flow", "How a request moves through the system")

    steps = [
        ("1", "GPS Input", "Raw coordinates from driver", BLUE),
        ("2", "Map Matching", "Snap to road segment", TEAL),
        ("3", "Demand Detection", "Check if service needed", ORANGE),
        ("4", "Candidate Search", "Find eligible stations", PURPLE),
        ("5", "Ranking", "Score and return best", RED),
    ]

    for i, (num, name, desc, color) in enumerate(steps):
        x = 0.5 + i * 1.9

        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.6), Inches(1.7), Inches(2))
        s.fill.solid()
        s.fill.fore_color.rgb = WHITE
        s.line.color.rgb = color
        s.line.width = Pt(2)

        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.6), Inches(1.75), Inches(0.5), Inches(0.5))
        n.fill.solid()
        n.fill.fore_color.rgb = color
        n.line.fill.background()
        nt = slide.shapes.add_textbox(Inches(x + 0.6), Inches(1.85), Inches(0.5), Inches(0.4))
        tf = nt.text_frame
        tf.paragraphs[0].text = num
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        name_t = slide.shapes.add_textbox(Inches(x + 0.1), Inches(2.35), Inches(1.5), Inches(0.4))
        tf = name_t.text_frame
        tf.paragraphs[0].text = name
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = color
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        desc_t = slide.shapes.add_textbox(Inches(x + 0.1), Inches(2.75), Inches(1.5), Inches(0.7))
        tf = desc_t.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = SLATE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        if i < 4:
            arrow_right(slide, x + 1.7, 2.5, 0.2)


# ============================================================
# SLIDE 4: Component Communication
# ============================================================
def slide_communication(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "Component Communication", "How services interact")

    # Center - FastAPI
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(2.3), Inches(3), Inches(1))
    c.fill.solid()
    c.fill.fore_color.rgb = NAVY
    c.line.fill.background()
    ct = slide.shapes.add_textbox(Inches(3.5), Inches(2.55), Inches(3), Inches(0.5))
    tf = ct.text_frame
    tf.paragraphs[0].text = "FastAPI Gateway"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Top - Client
    client = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(1.3), Inches(3), Inches(0.6))
    client.fill.solid()
    client.fill.fore_color.rgb = SLATE
    client.line.fill.background()
    ct = slide.shapes.add_textbox(Inches(3.5), Inches(1.4), Inches(3), Inches(0.4))
    tf = ct.text_frame
    tf.paragraphs[0].text = "Client"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    arrow_down(slide, 4.92, 1.95, 0.3)

    # Left - GraphHopper
    gh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(2.3), Inches(2.6), Inches(1))
    gh.fill.solid()
    gh.fill.fore_color.rgb = ORANGE
    gh.line.fill.background()
    ght = slide.shapes.add_textbox(Inches(0.4), Inches(2.55), Inches(2.6), Inches(0.5))
    tf = ght.text_frame
    tf.paragraphs[0].text = "GraphHopper"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    arrow_right(slide, 3.0, 2.7, 0.45)

    # Right - PostgreSQL
    pg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7), Inches(2.3), Inches(2.6), Inches(1))
    pg.fill.solid()
    pg.fill.fore_color.rgb = TEAL
    pg.line.fill.background()
    pgt = slide.shapes.add_textbox(Inches(7), Inches(2.55), Inches(2.6), Inches(0.5))
    tf = pgt.text_frame
    tf.paragraphs[0].text = "PostgreSQL"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    arrow_right(slide, 6.5, 2.7, 0.45)

    # Bottom - Redis
    redis1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4), Inches(2.6), Inches(0.8))
    redis1.fill.solid()
    redis1.fill.fore_color.rgb = BLUE
    redis1.line.fill.background()
    r1t = slide.shapes.add_textbox(Inches(1.5), Inches(4.15), Inches(2.6), Inches(0.5))
    tf = r1t.text_frame
    tf.paragraphs[0].text = "Redis (State)"
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    redis2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.9), Inches(4), Inches(2.6), Inches(0.8))
    redis2.fill.solid()
    redis2.fill.fore_color.rgb = BLUE
    redis2.line.fill.background()
    r2t = slide.shapes.add_textbox(Inches(5.9), Inches(4.15), Inches(2.6), Inches(0.5))
    tf = r2t.text_frame
    tf.paragraphs[0].text = "Redis (Cache)"
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    arrow_down(slide, 4.92, 3.3, 0.5)
    arrow_down(slide, 2.8, 3.3, 0.65)
    arrow_down(slide, 7.2, 3.3, 0.65)

    # Labels
    labels = [
        (4.2, 4.9, "Driver location"),
        (2.0, 4.9, "CAS updates"),
        (6.4, 4.9, "TTL snapshots"),
    ]
    for x, y, txt in labels:
        lt = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2), Inches(0.3))
        tf = lt.text_frame
        tf.paragraphs[0].text = txt
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = SLATE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER


# ============================================================
# SLIDE 5: Recommendation Workflow
# ============================================================
def slide_workflow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "Recommendation Workflow", "Orchestration")

    # Main workflow box
    wf = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3), Inches(1.4), Inches(4), Inches(0.7))
    wf.fill.solid()
    wf.fill.fore_color.rgb = NAVY
    wf.line.fill.background()
    wft = slide.shapes.add_textbox(Inches(3), Inches(1.55), Inches(4), Inches(0.4))
    tf = wft.text_frame
    tf.paragraphs[0].text = "RecommendationWorkflow"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    arrow_down(slide, 4.92, 2.1, 0.35)

    # Step 1 - Search
    s1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.6), Inches(4.3), Inches(1.3))
    s1.fill.solid()
    s1.fill.fore_color.rgb = WHITE
    s1.line.color.rgb = TEAL
    s1.line.width = Pt(2)

    s1t = slide.shapes.add_textbox(Inches(0.7), Inches(2.75), Inches(3.9), Inches(0.35))
    tf = s1t.text_frame
    tf.paragraphs[0].text = "1. Candidate Search"
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = TEAL

    s1_items = [
        "Resolve operational snapshots",
        "Route origin to station to destination",
        "Filter: OFFLINE, FULL, INCOMPATIBLE",
    ]
    for i, item in enumerate(s1_items):
        s1i = slide.shapes.add_textbox(Inches(0.7), Inches(3.15 + i * 0.28), Inches(3.9), Inches(0.28))
        tf = s1i.text_frame
        tf.paragraphs[0].text = f"  {item}"
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = SLATE

    # Step 2 - Rank
    s2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.2), Inches(2.6), Inches(4.3), Inches(1.3))
    s2.fill.solid()
    s2.fill.fore_color.rgb = WHITE
    s2.line.color.rgb = ORANGE
    s2.line.width = Pt(2)

    s2t = slide.shapes.add_textbox(Inches(5.4), Inches(2.75), Inches(3.9), Inches(0.35))
    tf = s2t.text_frame
    tf.paragraphs[0].text = "2. Ranking"
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE

    s2_items = [
        "TOTAL_SERVICE_COMPLETION_V1",
        "Cost = travel + queue + service + detour",
        "Return: best + top-N alternatives",
    ]
    for i, item in enumerate(s2_items):
        s2i = slide.shapes.add_textbox(Inches(5.4), Inches(3.15 + i * 0.28), Inches(3.9), Inches(0.28))
        tf = s2i.text_frame
        tf.paragraphs[0].text = f"  {item}"
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = SLATE

    arrow_right(slide, 4.85, 3.15, 0.3)

    # Conflict handling
    cf = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(4.2), Inches(9), Inches(1.1))
    cf.fill.solid()
    cf.fill.fore_color.rgb = RgbColor(0xFE, 0xF2, 0xF2)
    cf.line.color.rgb = RED
    cf.line.width = Pt(1)

    cft = slide.shapes.add_textbox(Inches(0.7), Inches(4.35), Inches(8.6), Inches(0.35))
    tf = cft.text_frame
    tf.paragraphs[0].text = "Conflict Handling: CandidateStateChanged"
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RED

    cf_desc = slide.shapes.add_textbox(Inches(0.7), Inches(4.75), Inches(8.6), Inches(0.45))
    tf = cf_desc.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "Queue surge during search -> retry once -> if conflict persists, return HTTP 409"
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = SLATE


# ============================================================
# SLIDE 6: Snapshot Architecture
# ============================================================
def slide_snapshots(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "Snapshot Architecture", "Redis cache + PostgreSQL authority")

    # Three snapshot types
    types = [
        (0.4, "Station", "status, services, capacity", TEAL),
        (3.55, "Queue", "wait_time, queue_length", ORANGE),
        (6.7, "Traffic", "congestion polygons", PURPLE),
    ]
    for x, name, desc, color in types:
        t = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.4), Inches(2.95), Inches(0.9))
        t.fill.solid()
        t.fill.fore_color.rgb = color
        t.line.fill.background()
        tt = slide.shapes.add_textbox(Inches(x), Inches(1.5), Inches(2.95), Inches(0.35))
        tf = tt.text_frame
        tf.paragraphs[0].text = name
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        td = slide.shapes.add_textbox(Inches(x), Inches(1.85), Inches(2.95), Inches(0.35))
        tf = td.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Resolution flow
    flow_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(2.6), Inches(9.2), Inches(1.5))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = WHITE
    flow_box.line.color.rgb = RgbColor(0xE2, 0xE8, 0xF0)
    flow_box.line.width = Pt(1)

    flow_title = slide.shapes.add_textbox(Inches(0.6), Inches(2.7), Inches(8.8), Inches(0.35))
    tf = flow_title.text_frame
    tf.paragraphs[0].text = "SnapshotResolver Flow"
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = NAVY

    flow_steps = [
        (0.8, "1. Redis Cache", "TTL-based"),
        (3.5, "2. PostgreSQL", "Persistent"),
        (6.2, "3. Return", "To caller"),
    ]
    for x, name, note in flow_steps:
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.1), Inches(2.4), Inches(0.8))
        s.fill.solid()
        s.fill.fore_color.rgb = BLUE
        s.line.fill.background()
        st = slide.shapes.add_textbox(Inches(x), Inches(3.2), Inches(2.4), Inches(0.4))
        tf = st.text_frame
        tf.paragraphs[0].text = name
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        sn = slide.shapes.add_textbox(Inches(x), Inches(3.55), Inches(2.4), Inches(0.3))
        tf = sn.text_frame
        tf.paragraphs[0].text = note
        tf.paragraphs[0].font.size = Pt(9)
        tf.paragraphs[0].font.color.rgb = RgbColor(0xB4, 0xC4, 0xD4)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        if x < 6:
            arrow_right(slide, x + 2.45, 3.4, 0.25)

    # Ingestion
    ing_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(4.3), Inches(9.2), Inches(1))
    ing_box.fill.solid()
    ing_box.fill.fore_color.rgb = WHITE
    ing_box.line.color.rgb = RgbColor(0xE2, 0xE8, 0xF0)
    ing_box.line.width = Pt(1)

    ing_title = slide.shapes.add_textbox(Inches(0.6), Inches(4.4), Inches(8.8), Inches(0.35))
    tf = ing_title.text_frame
    tf.paragraphs[0].text = "Ingestion (Internal)"
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RED

    ing_desc = slide.shapes.add_textbox(Inches(0.6), Inches(4.8), Inches(8.8), Inches(0.4))
    tf = ing_desc.text_frame
    tf.paragraphs[0].text = "Simulator.tick() -> IngestionService -> PostgreSQL"
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = SLATE


# ============================================================
# SLIDE 7: Key Design Rules
# ============================================================
def slide_design(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "Key Design Rules", "Architectural decisions")

    rules = [
        ("1", "No Business Logic in Frontend", "UI calls actual API endpoints. Calculations stay in backend.", BLUE),
        ("2", "PostgreSQL = Snapshot Authority", "Redis is optional cache. PostgreSQL holds immutable truth.", TEAL),
        ("3", "Conflict = 409 + Retry", "Queue surge -> CandidateStateChanged -> retry search once.", ORANGE),
        ("4", "GPS State with CAS", "Per-driver Redis updates with version-based conflict detection.", PURPLE),
        ("5", "Shared Routing Adapter", "Single httpx client shared across all routing calls.", RED),
    ]

    for i, (num, title, desc, color) in enumerate(rules):
        y = 1.4 + i * 0.82

        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(y + 0.1), Inches(0.5), Inches(0.5))
        n.fill.solid()
        n.fill.fore_color.rgb = color
        n.line.fill.background()
        nt = slide.shapes.add_textbox(Inches(0.5), Inches(y + 0.18), Inches(0.5), Inches(0.4))
        tf = nt.text_frame
        tf.paragraphs[0].text = num
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        t = slide.shapes.add_textbox(Inches(1.2), Inches(y + 0.05), Inches(8.3), Inches(0.35))
        tf = t.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = color

        d = slide.shapes.add_textbox(Inches(1.2), Inches(y + 0.4), Inches(8.3), Inches(0.35))
        tf = d.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.color.rgb = SLATE


# ============================================================
# SLIDE 8: Lifecycle
# ============================================================
def slide_lifecycle(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_title_bar(slide, prs, "Application Lifecycle", "Startup -> Runtime -> Shutdown")

    # Startup
    start_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(1.5), Inches(2.9), Inches(3.5))
    start_box.fill.solid()
    start_box.fill.fore_color.rgb = WHITE
    start_box.line.color.rgb = TEAL
    start_box.line.width = Pt(2)

    start_title = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(2.5), Inches(0.4))
    tf = start_title.text_frame
    tf.paragraphs[0].text = "Startup"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = TEAL

    start_items = [
        "httpx.AsyncClient",
        "asyncpg.create_pool",
        "Redis.from_url",
        "SnapshotResolver",
        "RecommendationWorkflow",
        "RealtimeSimulator",
    ]
    for i, item in enumerate(start_items):
        itb = slide.shapes.add_textbox(Inches(0.6), Inches(2.1 + i * 0.4), Inches(2.5), Inches(0.4))
        tf = itb.text_frame
        tf.paragraphs[0].text = f"  {item}"
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.color.rgb = SLATE

    # Runtime
    runtime_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.55), Inches(1.5), Inches(2.9), Inches(3.5))
    runtime_box.fill.solid()
    runtime_box.fill.fore_color.rgb = WHITE
    runtime_box.line.color.rgb = BLUE
    runtime_box.line.width = Pt(2)

    runtime_title = slide.shapes.add_textbox(Inches(3.75), Inches(1.6), Inches(2.5), Inches(0.4))
    tf = runtime_title.text_frame
    tf.paragraphs[0].text = "Runtime"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE

    runtime_items = [
        "Request orchestration",
        "Candidate search",
        "Ranking",
        "Snapshot caching",
        "GPS state updates",
        "Background simulation",
    ]
    for i, item in enumerate(runtime_items):
        itb = slide.shapes.add_textbox(Inches(3.75), Inches(2.1 + i * 0.4), Inches(2.5), Inches(0.4))
        tf = itb.text_frame
        tf.paragraphs[0].text = f"  {item}"
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.color.rgb = SLATE

    # Shutdown
    shutdown_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.5), Inches(2.9), Inches(3.5))
    shutdown_box.fill.solid()
    shutdown_box.fill.fore_color.rgb = WHITE
    shutdown_box.line.color.rgb = RED
    shutdown_box.line.width = Pt(2)

    shutdown_title = slide.shapes.add_textbox(Inches(6.9), Inches(1.6), Inches(2.5), Inches(0.4))
    tf = shutdown_title.text_frame
    tf.paragraphs[0].text = "Shutdown"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RED

    shutdown_items = [
        "Simulator.stop()",
        "Clear app.state",
        "Close httpx client",
        "Close segment resolver",
    ]
    for i, item in enumerate(shutdown_items):
        itb = slide.shapes.add_textbox(Inches(6.9), Inches(2.1 + i * 0.4), Inches(2.5), Inches(0.4))
        tf = itb.text_frame
        tf.paragraphs[0].text = f"  {item}"
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.color.rgb = SLATE

    # Arrows
    arrow_right(slide, 3.3, 3.1, 0.2)
    arrow_right(slide, 6.45, 3.1, 0.2)


# ============================================================
# SLIDE 9: End
# ============================================================
def slide_end(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()

    dec = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(-2), Inches(6), Inches(6))
    dec.fill.solid()
    dec.fill.fore_color.rgb = RgbColor(0x2D, 0x4A, 0x6F)
    dec.line.fill.background()

    t = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.5))
    tf = t.text_frame
    tf.paragraphs[0].text = "Questions?"
    tf.paragraphs[0].font.size = Pt(56)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    d = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(0.8))
    tf = d.text_frame
    tf.paragraphs[0].text = "Demo: /demo"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    slide_title(prs)
    slide_overview(prs)
    slide_flow(prs)
    slide_communication(prs)
    slide_workflow(prs)
    slide_snapshots(prs)
    slide_design(prs)
    slide_lifecycle(prs)
    slide_end(prs)

    output = "e:\\build6week\\EV_Arch_Mentor_v4.pptx"
    prs.save(output)
    print(f"[OK] Created: {output}")
    print(f"[INFO] Total: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
