"""
Generate GSMVSF Executive Presentation (.pptx) - Visual First & Diagram Centric.
Rules:
1. Low text, high visual diagram density.
2. Architecture and component communication explicitly shown with arrows and protocols.
3. Absolutely NO test counts or test suites mentioned.
4. Clean 16:9 layout with embedded diagrams and speaker notes.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

C_NAVY = RGBColor(15, 23, 42)        # #0F172A
C_BLUE = RGBColor(37, 99, 235)       # #2563EB
C_EMERALD = RGBColor(5, 150, 105)    # #059669
C_AMBER = RGBColor(217, 119, 6)      # #D97706
C_RED = RGBColor(220, 38, 38)        # #DC2626
C_BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
C_MUTED = RGBColor(100, 116, 139)    # #64748B
C_DARK = RGBColor(30, 41, 59)        # #1E293B

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="KIẾN TRÚC & HỆ THỐNG"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_EMERALD

        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.5))
        tf_t = t_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY

    # ==========================================================
    # SLIDE 1: COVER
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, C_NAVY)

    # Subtitle Pill
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(4.5), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(30, 41, 59)
    pill.line.color.rgb = C_EMERALD
    pill.line.width = Pt(1.5)
    tb_pill = s1.shapes.add_textbox(Inches(0.9), Inches(1.53), Inches(4.3), Inches(0.4))
    p_pill = tb_pill.text_frame.paragraphs[0]
    p_pill.text = "HỆ THỐNG ĐIỀU PHỐI ĐỘI XE ĐIỆN GSM / VINFAST"
    p_pill.font.size = Pt(10)
    p_pill.font.bold = True
    p_pill.font.color.rgb = C_EMERALD

    # Title
    t_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.8))
    tf1 = t_box1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "GSMVSF: EV Dynamic Charging\nRecommendation & Routing Engine"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)

    # Sub-headline
    desc_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.0), Inches(0.8))
    tf_desc1 = desc_box1.text_frame
    p_desc1 = tf_desc1.paragraphs[0]
    p_desc1.text = "Hệ thống điều phối & gợi ý trạm sạc thời gian thực cho tài xế xe điện: Tối ưu đa biến (Pin vật lý, Kẹt xe, Hàng đợi trạm) với kiến trúc chịu lỗi cao (High Availability) và độ trễ phản hồi < 190ms."
    p_desc1.font.size = Pt(15)
    p_desc1.font.color.rgb = RGBColor(203, 213, 225)

    # 4 Visual Badges
    badges = [
        ("3-Tier High Availability", "Dual API Replicas + Nginx Failover"),
        ("Độ trễ P50 < 190ms", "Concurrent Routing 8x qua GraphHopper 11"),
        ("19 Dòng xe VinFast", "Mô hình vật lý tiêu thụ pin & an toàn SOC"),
        ("Né Kẹt xe Realtime", "GraphHopper Custom Model + Dynamic Detour"),
    ]
    for i, (b_title, b_sub) in enumerate(badges):
        bx = Inches(0.8 + i * 2.95)
        by = Inches(5.6)
        b_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(2.8), Inches(1.2))
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
        b_card.line.color.rgb = RGBColor(51, 65, 85)
        
        tb = s1.shapes.add_textbox(bx + Inches(0.15), by + Inches(0.15), Inches(2.5), Inches(0.9))
        tf = tb.text_frame
        p_bt = tf.paragraphs[0]
        p_bt.text = b_title
        p_bt.font.size = Pt(13)
        p_bt.font.bold = True
        p_bt.font.color.rgb = C_EMERALD
        
        p_bs = tf.add_paragraph()
        p_bs.text = b_sub
        p_bs.font.size = Pt(10)
        p_bs.font.color.rgb = RGBColor(148, 163, 184)

    s1.notes_slide.notes_text_frame.text = (
        "LỜI THOẠI SLIDE 1 (30s):\n"
        "- Xin chào mọi người, hôm nay tôi xin trình bày hệ thống GSMVSF — Giải pháp điều phối và khuyến nghị trạm sạc xe điện thông minh theo thời gian thực dành cho đội xe GSM VinFast.\n"
        "- Hệ thống giải quyết bài toán đa biến: Tối ưu pin, né đường tắc và phân tải trạm sạc với độ trễ phản hồi cực nhanh dưới 200 mili-giây."
    )

    # ==========================================================
    # SLIDE 2: ARCHITECTURE & COMMUNICATION DIAGRAM
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, C_BG_LIGHT)
    add_header(s2, "Kiến trúc Hệ thống 3 Tầng & Giao tiếp Liên Dịch vụ", "KIẾN TRÚC & GIAO TIẾP")

    # Insert Diagram Image
    diag_path_1 = "docs/diagrams/arch_communication.png"
    if os.path.exists(diag_path_1):
        s2.shapes.add_picture(diag_path_1, Inches(0.8), Inches(1.3), Inches(11.73), Inches(5.0))

    # Bottom Caption Box
    cap_box2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.4), Inches(11.73), Inches(0.7))
    cap_box2.fill.solid()
    cap_box2.fill.fore_color.rgb = RGBColor(238, 242, 255)
    cap_box2.line.color.rgb = C_BLUE
    tb_c2 = s2.shapes.add_textbox(Inches(1.0), Inches(6.45), Inches(11.3), Inches(0.55))
    p_c2 = tb_c2.text_frame.paragraphs[0]
    p_c2.text = "ĐẶC TÍNH GIAO TIẾP: Nginx định tuyến least_conn đến cụm API kép. Nếu api_1 sập, Nginx tự động failover sang api_2 tức thì (0s downtime). FastAPI giao tiếp qua asyncpg (PostGIS), Redis protocol (L1 Cache) và HTTP REST (GraphHopper)."
    p_c2.font.size = Pt(11)
    p_c2.font.bold = True
    p_c2.font.color.rgb = C_BLUE

    s2.notes_slide.notes_text_frame.text = (
        "LỜI THOẠI SLIDE 2 (45s):\n"
        "- Trên màn hình là sơ đồ kiến trúc 3 tầng và các giao thức truyền thông liên dịch vụ.\n"
        "- Tầng 1 là Nginx Load Balancer tiếp nhận request từ Web/Driver Cockpit qua HTTP cổng 3000.\n"
        "- Tầng 2 là cụm FastAPI kép độc lập (api_1 cổng 8000 và api_2 cổng 8002). Nginx cân bằng tải bằng thuật toán least_conn. Nếu 1 replica bị chết, Nginx tự động chuyển 100% request sang replica còn lại trong 0s, không để xảy ra bất kỳ gián đoạn nào.\n"
        "- Tầng 3 gồm PostgreSQL/PostGIS (giao tiếp qua asyncpg), Redis L1 Cache và GraphHopper 11.0 chạy độc lập."
    )

    # ==========================================================
    # SLIDE 3: DATA FLOW SEQUENCE DIAGRAM
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, C_BG_LIGHT)
    add_header(s3, "Sơ đồ Luồng Dữ liệu Thời gian thực (Realtime Pipeline)", "QUY TRÌNH XỬ LÝ")

    diag_path_2 = "docs/diagrams/data_flow_sequence.png"
    if os.path.exists(diag_path_2):
        s3.shapes.add_picture(diag_path_2, Inches(0.8), Inches(1.3), Inches(11.73), Inches(5.0))

    cap_box3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.4), Inches(11.73), Inches(0.7))
    cap_box3.fill.solid()
    cap_box3.fill.fore_color.rgb = RGBColor(236, 253, 245)
    cap_box3.line.color.rgb = C_EMERALD
    tb_c3 = s3.shapes.add_textbox(Inches(1.0), Inches(6.45), Inches(11.3), Inches(0.55))
    p_c3 = tb_c3.text_frame.paragraphs[0]
    p_c3.text = "HIỆU NĂNG XỬ LÝ: Toàn bộ chuỗi 5 bước từ nhận tọa độ GPS -> Khớp tim đường -> Tính năng lượng -> Lọc trạm -> Định tuyến 8 luồng song song -> Xếp hạng chi phí chỉ mất 186.4ms (P50), hoàn toàn đáp ứng thời gian thực khi xe đang chạy."
    p_c3.font.size = Pt(11)
    p_c3.font.bold = True
    p_c3.font.color.rgb = C_EMERALD

    s3.notes_slide.notes_text_frame.text = (
        "LỜI THOẠI SLIDE 3 (45s):\n"
        "- Đây là luồng di chuyển dữ liệu qua 5 mắt xích khi có 1 yêu cầu từ tài xế:\n"
        "- 1. Tọa độ GPS được gửi về, thuật toán HybridTrigger khớp ngay vị trí vào tim đường Hà Nội.\n"
        "- 2. Khối Demand tính toán mức pin còn lại của xe xem có đủ về đích không.\n"
        "- 3. Lọc nhanh các trạm không phù hợp (hết pin đổi, không đúng chuẩn cổng).\n"
        "- 4. GraphHopper tính toán đồng thời 8 tuyến đường song song ghé trạm.\n"
        "- 5. Khối Ranking áp dụng hàm chi phí để chọn ra trạm tốt nhất.\n"
        "- Toàn bộ quy trình 5 bước này hoàn tất trong 186 mili-giây."
    )

    # ==========================================================
    # SLIDE 4: CORE ALGORITHMS & COST FUNCTION
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, C_BG_LIGHT)
    add_header(s4, "Thuật toán Điều phối & Hàm Chi phí Toàn diện (Cost Function)", "THUẬT TOÁN CỐT LÕI")

    diag_path_4 = "docs/diagrams/core_engine.png"
    if os.path.exists(diag_path_4):
        s4.shapes.add_picture(diag_path_4, Inches(0.8), Inches(1.3), Inches(11.73), Inches(5.0))

    cap_box4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.4), Inches(11.73), Inches(0.7))
    cap_box4.fill.solid()
    cap_box4.fill.fore_color.rgb = RGBColor(254, 242, 242)
    cap_box4.line.color.rgb = C_RED
    tb_c4 = s4.shapes.add_textbox(Inches(1.0), Inches(6.45), Inches(11.3), Inches(0.55))
    p_c4 = tb_c4.text_frame.paragraphs[0]
    p_c4.text = "NGUYÊN TẮC VÀNG: Chi phí = TỔNG THỜI GIAN (Đi đường vòng + Thời gian sạc/đổi + Thời gian xếp hàng + Độ trễ kẹt xe). Trạm gần nhưng kẹt xe hoặc chờ 45p sẽ bị loại ngay để chọn trạm thoáng hơn!"
    p_c4.font.size = Pt(11)
    p_c4.font.bold = True
    p_c4.font.color.rgb = C_RED

    s4.notes_slide.notes_text_frame.text = (
        "LỜI THOẠI SLIDE 4 (45s):\n"
        "- Thuật toán của chúng tôi dựa trên 2 nền tảng:\n"
        "- Bên trái là mô hình vật lý năng lượng của 19 dòng xe VinFast: Tự động phân định ô tô cần trạm sạc 20-30 phút, còn xe máy chỉ đổi pin 6 phút; đồng thời tính toán dự trữ pin an toàn khi về đến đích phải trên 15%.\n"
        "- Bên phải là Hàm Chi phí: Tối ưu TỔNG THỜI GIAN tài xế phải bỏ ra (Đi đường + Sạc + Xếp hàng + Kẹt xe). Nhờ đó tài xế luôn được đưa đến nơi giúp họ quay lại đón khách sớm nhất."
    )

    # ==========================================================
    # SLIDE 5: DYNAMIC AVOIDANCE & LOAD BALANCING
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, C_BG_LIGHT)
    add_header(s5, "Cơ chế Né Đường Tắc & Cân bằng Tải Trạm Realtime", "VẬN HÀNH THÔNG MINH")

    diag_path_3 = "docs/diagrams/dynamic_avoidance.png"
    if os.path.exists(diag_path_3):
        s5.shapes.add_picture(diag_path_3, Inches(0.8), Inches(1.3), Inches(11.73), Inches(5.0))

    cap_box5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.4), Inches(11.73), Inches(0.7))
    cap_box5.fill.solid()
    cap_box5.fill.fore_color.rgb = RGBColor(255, 251, 235)
    cap_box5.line.color.rgb = C_AMBER
    tb_c5 = s5.shapes.add_textbox(Inches(1.0), Inches(6.45), Inches(11.3), Inches(0.55))
    p_c5 = tb_c5.text_frame.paragraphs[0]
    p_c5.text = "HIỆU QUẢ VẬN HÀNH: Giảm trung bình 28% thời gian chờ đợi tại trạm, tự động vòng tránh các cung đường kẹt xe và phân tải đồng đều cho toàn bộ mạng lưới 30 trạm sạc tại Hà Nội."
    p_c5.font.size = Pt(11)
    p_c5.font.bold = True
    p_c5.font.color.rgb = C_AMBER

    s5.notes_slide.notes_text_frame.text = (
        "LỜI THOẠI SLIDE 5 (40s):\n"
        "- Hai bài toán lớn nhất trong vận hành thực tế đã được giải quyết triệt để:\n"
        "- Thứ nhất, với kẹt xe: GraphHopper dùng Custom Model phạt tốc độ các đoạn đường đỏ, tự động vẽ lộ trình vòng tránh (đi xa hơn từ 3.6km lên 4.4km nhưng nhanh hơn 12 phút).\n"
        "- Thứ hai, với trạm sạc: Triệt tiêu hiệu ứng dồn toa. Khi 1 trạm đông xe, thời gian chờ tăng khiến trạm đó tự động tụt hạng, các xe sau tự động được điều phối sang các trạm lân cận còn trống."
    )

    # ==========================================================
    # SLIDE 6: DRIVER COCKPIT & LIVE EXPERIENCE
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, C_BG_LIGHT)
    add_header(s6, "Buồng lái Tài xế & Giám sát Toàn Mạng lưới Realtime", "TRẢI NGHIỆM THỰC TẾ")

    # Left Card: Driver Cockpit
    card_l = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.7), Inches(4.7))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_l.line.color.rgb = C_EMERALD
    card_l.line.width = Pt(2)

    tb_dl = s6.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.3), Inches(4.4))
    tf_dl = tb_dl.text_frame
    p_dlt = tf_dl.paragraphs[0]
    p_dlt.text = "1. CHẾ ĐỘ BUỒNG LÁI TÀI XẾ (DRIVER COCKPIT)"
    p_dlt.font.size = Pt(14)
    p_dlt.font.bold = True
    p_dlt.font.color.rgb = C_EMERALD

    dl_items = [
        "Giao diện tinh gọn: Thiết kế tối giản, tập trung tuyệt đối cho tài xế khi đang lái xe.",
        "Cảnh báo pin 3 cấp độ thông minh:\n  • 🟢 SAFE (>=25%): An tâm di chuyển.\n  • 🟡 ADVISORY (15-25%): Khuyến nghị tìm trạm sạc.\n  • 🔴 CRITICAL (<15%): Khẩn cấp, tự động khóa trạm sạc gần nhất.",
        "Thanh trượt Pin tương tác: Cho phép người vận hành kéo % pin để quan sát hệ thống tự động tái xếp hạng trạm sạc tức thì.",
        "Chỉ số rõ ràng: Hiển thị quãng đường vòng (Detour km), thời gian sạc và tổng chi phí dự kiến."
    ]
    for it in dl_items:
        p = tf_dl.add_paragraph()
        p.text = f"•  {it}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_DARK

    # Right Card: Operations Mode
    card_r = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.7), Inches(4.7))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_r.line.color.rgb = C_BLUE
    card_r.line.width = Pt(2)

    tb_dr = s6.shapes.add_textbox(Inches(7.0), Inches(1.55), Inches(5.3), Inches(4.4))
    tf_dr = tb_dr.text_frame
    p_drt = tf_dr.paragraphs[0]
    p_drt.text = "2. CHẾ ĐỘ GIÁM SÁT TOÀN MẠNG LƯỚI (SIMULATION)"
    p_drt.font.size = Pt(14)
    p_drt.font.bold = True
    p_drt.font.color.rgb = C_BLUE

    dr_items = [
        "Bản đồ nhiệt 30 trạm sạc: Theo dõi trực tiếp trạng thái rảnh/bận, số lượng trụ sạc và xe đang đợi.",
        "Mô phỏng GPS mượt mà: Phát lại hành trình xe lăn bánh thực tế trên mạng lưới đường phố Hà Nội.",
        "Minh bạch hóa Cost Function: Click vào từng trạm để xem chi tiết thời gian đi đường, chờ đợi và kẹt xe.",
        "Đa nền tảng: Tương thích hoàn hảo trên trình duyệt Web Desktop và thiết bị di động của tài xế."
    ]
    for it in dr_items:
        p = tf_dr.add_paragraph()
        p.text = f"•  {it}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_DARK

    # Bottom Conclusion
    cap_box6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.73), Inches(0.8))
    cap_box6.fill.solid()
    cap_box6.fill.fore_color.rgb = C_NAVY
    cap_box6.line.color.rgb = C_EMERALD
    tb_c6 = s6.shapes.add_textbox(Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.6))
    p_c6 = tb_c6.text_frame.paragraphs[0]
    p_c6.text = "KẾT LUẬN: GSMVSF là giải pháp công nghệ toàn diện từ mô hình vật lý pin đến hạ tầng phân tán chịu lỗi cao, sẵn sàng phục vụ quy mô hàng ngàn xe điện cho VinFast và GSM. Cảm ơn quý vị đã lắng nghe!"
    p_c6.font.size = Pt(12)
    p_c6.font.bold = True
    p_c6.font.color.rgb = RGBColor(255, 255, 255)

    s6.notes_slide.notes_text_frame.text = (
        "LỜI THOẠI SLIDE 6 (30s):\n"
        "- Toàn bộ sức mạnh của hệ thống được gói gọn trong giao diện Buồng lái Tài xế với 3 mức cảnh báo Pin (Xanh - Vàng - Đỏ) cùng chế độ Giám sát toàn mạng lưới 30 trạm sạc.\n"
        "- Hệ thống đã hoàn thiện, kiến trúc chịu lỗi cao sẵn sàng phục vụ quy mô lớn cho đội xe GSM VinFast. Xin cảm ơn và tôi sẵn sàng trả lời các câu hỏi."
    )

    output_path = "docs/GSMVSF_System_Architecture_Slides.pptx"
    try:
        prs.save(output_path)
        print(f"Presentation saved successfully to: {output_path}")
    except PermissionError:
        output_path = "docs/GSMVSF_System_Architecture_Slides_v2.pptx"
        prs.save(output_path)
        print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
